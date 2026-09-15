# 生成给队友A的 Excel 填表模板（防呆版：下拉选择，不用手打）
# 阮不用看懂这个脚本，跑一次就生成 data/数据登记表_模板.xlsx
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = r"C:/Users/Lenovo/WorkBuddy/黑客松/loreal-hackathon-beautyproof/data/数据登记表_模板.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "填这里"

headers = ["file_name", "label", "tamper_type", "tamper_region", "source", "note"]

# ---- 第1行：表头（黄底加粗，提醒别动）----
head_fill = PatternFill("solid", fgColor="FFF2CC")
head_font = Font(bold=True, size=12)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for c, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.fill = head_fill
    cell.font = head_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border

# ---- 第2-7行：填好的样例（浅灰底，作为照抄参考）----
samples = [
    ["001_real_lipstick.jpg", "1", "none", "none", "小红书", "带品牌水印"],
    ["002_real_foundation.jpg", "1", "none", "none", "微博", "有价格文字"],
    ["003_fake_p图.jpg", "0", "p图", "嘴唇区域改色", "自造", "豆沙色改成正红"],
    ["004_fake_拼接.jpg", "0", "拼接", "右侧口红是从另一张图抠来的", "自造", "抠图边缘有点糊"],
    ["005_fake_ai生成.jpg", "0", "ai生成", "整张图都是AI生成", "自造", "即梦生成"],
    ["006_fake_改文字.jpg", "0", "改文字", "左下角价格标签", "自造", "199改成99"],
]
gray = PatternFill("solid", fgColor="F2F2F2")
for r, row in enumerate(samples, start=2):
    for c, v in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.fill = gray
        cell.border = border
        cell.alignment = Alignment(vertical="center")

# ---- 第8行起：空白，队友从这里开始填 ----
for r in range(8, 208):
    for c in range(1, 7):
        ws.cell(row=r, column=c).border = border

# ---- 下拉选择（防呆：点一下就能选，不用打字）----
dv_label = DataValidation(type="list", formula1='"1,0"', allow_blank=True, showDropDown=False)
dv_type = DataValidation(type="list", formula1='"none,p图,拼接,ai生成,改文字"', allow_blank=True, showDropDown=False)
dv_source = DataValidation(type="list", formula1='"小红书,微博,抖音,自造"', allow_blank=True, showDropDown=False)
for dv in (dv_label, dv_type, dv_source):
    ws.add_data_validation(dv)
dv_label.add("B8:B207")
dv_type.add("C8:C207")
dv_source.add("E8:E207")

# ---- 列宽 ----
for col, w in zip("ABCDEF", [30, 8, 14, 34, 12, 26]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"  # 冻结第一行，往下翻也能看到表头

# ---- 第二个表：说明 ----
ws2 = wb.create_sheet("先看这个")
notes = [
    ["怎么用这个表", ""],
    ["", ""],
    ["第1步", "下载这个 Excel 文件，双击用 Excel（或 WPS）打开"],
    ["第2步", "点下面的「填这里」标签页"],
    ["第3步", "看到前 6 行灰色的是例子，照着它们的样子填"],
    ["第4步", "从**第 8 行**开始填你自己整理的第一张图"],
    ["第5步", "填完直接保存，把这个文件发回给阮"],
    ["", ""],
    ["每一列填什么（大白话）", ""],
    ["file_name", "图片文件名。格式：编号_real或fake_类别.jpg，比如 003_fake_p图.jpg"],
    ["label", "这张图是真的还是假的？点一下格子右边的小箭头选：1=真图，0=假图"],
    ["tamper_type", "如果是假图，属于哪一类？点箭头选：p图 / 拼接 / ai生成 / 改文字。真图选 none"],
    ["tamper_region", "改在哪个地方？用大白话写，比如「左下角价格标签」。真图填 none"],
    ["source", "图片哪来的？点箭头选：小红书 / 微博 / 抖音 / 自造"],
    ["note", "备注，想写啥写啥，不写就空着"],
    ["", ""],
    ["⚠️ 三件千万别做的事", ""],
    ["1", "别删第一行，也别改第一行的字（那是机器认的标签，改了就白填）"],
    ["2", "别把第一行六个词的顺序调换"],
    ["3", "别把文件另存成 CSV。就存成 Excel 原来的 .xlsx 格式发回来就行"],
]
for r, (a, b) in enumerate(notes, start=1):
    ws2.cell(row=r, column=1, value=a).font = Font(bold=True)
    ws2.cell(row=r, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
ws2.column_dimensions["A"].width = 26
ws2.column_dimensions["B"].width = 78

wb.save(OUT)
print("OK ->", OUT)
